"""Report writers: an xlsx workbook and a self-contained HTML dashboard."""

from __future__ import annotations

import html
from pathlib import Path

REPORT_COLUMNS = [
    "Provider", "Internal Ref", "Provider Ref", "Value Date",
    "Amount (Internal)", "Amount (Provider)", "Difference", "Currency",
    "Side", "Status", "Type", "Method", "Cleared Against", "Description",
]


def write_xlsx(result: dict, path: Path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3A4A3F")

    run_date = result["run_date"]
    cols = ["Run Date"] + REPORT_COLUMNS

    def sheet(title, records):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(horizontal="left")
        for r in records:
            d = r.as_dict()
            ws.append([run_date] + [d[c] for c in REPORT_COLUMNS])
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols, 1):
            cell_lens = [len(str(run_date))] if col == "Run Date" else \
                [len(str(r.as_dict().get(col, ""))) for r in records]
            width = max([len(col)] + cell_lens)
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(width + 3, 60)
        return ws

    # summary sheet first
    summ = wb.active
    summ.title = "Summary"
    summ.append(["Recon Daily — run", result["run_date"]])
    summ.append(["Look-back window (days)", result["lookback_days"]])
    summ.append(["Prior days available", ", ".join(result["prior_dates"]) or "none"])
    summ.append([])
    summ.append(["Internal rows", result["internal_count"]])
    summ.append(["Matched (same day)", len(result["matched"])])
    summ.append(["Cleared (aged)", len(result["cleared"])])
    summ.append(["Exceptions (manual)", len(result["exceptions"])])
    summ.append([])
    summ.append(["Provider", "Matched", "Cleared", "Exceptions", "Total"])
    for c in summ[10]:
        c.font = head_font
        c.fill = head_fill
    for p in result["per_provider"]:
        summ.append([p["provider"], p["matched"], p["cleared"], p["exceptions"],
                     p["matched"] + p["cleared"] + p["exceptions"]])
    summ.column_dimensions["A"].width = 24
    for col in "BCDE":
        summ.column_dimensions[col].width = 14

    sheet("Exceptions", result["exceptions"])
    sheet("Cleared", result["cleared"])
    sheet("Matched", result["matched"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_daily_snapshot(result: dict, lsum: dict, path: Path):
    """Dated record for the date folder: recon summary + the open-exceptions
    state as of this run + what was newly opened and cleared today."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from .ledger import OPEN_COLS, CLEARED_COLS

    wb = openpyxl.Workbook()
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3A4A3F")

    summ = wb.active
    summ.title = "Summary"
    for label, val in [
        ("Recon Daily — run date", result["run_date"]),
        ("Look-back window (days)", result["lookback_days"]),
        ("Prior days available", ", ".join(result["prior_dates"]) or "none"),
        ("", ""),
        ("Internal rows (today)", result["internal_count"]),
        ("Matched (same day)", len(result["matched"])),
        ("Cleared vs prior day (timing)", len(result["cleared"])),
        ("New exceptions opened today", lsum["new_today"]),
        ("Exceptions cleared today", lsum["cleared_today"]),
        ("Open exceptions (total)", lsum["open_total"]),
        ("  of which Manual Review (>%dd)" % result["lookback_days"], lsum["open_manual"]),
    ]:
        summ.append([label, val])
    summ.append([])
    hdr_row = summ.max_row + 1
    summ.append(["Provider", "Matched", "Cleared", "Exceptions", "Total"])
    for c in summ[hdr_row]:
        c.font = hf
        c.fill = fill
    for p in result["per_provider"]:
        summ.append([p["provider"], p["matched"], p["cleared"], p["exceptions"],
                     p["matched"] + p["cleared"] + p["exceptions"]])
    summ.column_dimensions["A"].width = 34
    for col in "BCDE":
        summ.column_dimensions[col].width = 14

    def sheet(title, cols, rows):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font = hf
            c.fill = fill
        for row in rows:
            ws.append([row.get(c, "") for c in cols])
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols, 1):
            lens = [len(col)] + [len(str(row.get(col, ""))) for row in rows]
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(lens) + 3, 60)

    sheet("Open Exceptions", OPEN_COLS, lsum["open_rows"])
    sheet("New Today", OPEN_COLS, lsum["new_today_rows"])
    sheet("Cleared Today", CLEARED_COLS, lsum["cleared_today_rows"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------
# HTML dashboard — self-contained, styled to match the browser tool
# ---------------------------------------------------------------------------
_CSS = """
:root{--ink:#12130f;--paper:#faf9f5;--line:#20221c;--slate:#3a4a3f;--moss:#5a7a52;
--clay:#b5533a;--amber:#c8892a;--grid:#e6e3da;--muted:#6b6d63;--ok:#4a6b3f;--warn:#b5533a;--info:#3a4a6b;
--mono:ui-monospace,'SF Mono',Menlo,Consolas,monospace;--sans:'Helvetica Neue',Arial,sans-serif;}
*{box-sizing:border-box;margin:0;padding:0}
body{background:var(--paper);color:var(--ink);font-family:var(--sans);line-height:1.5}
.wrap{max-width:1180px;margin:0 auto;padding:32px 24px 80px}
header{border-bottom:2px solid var(--line);padding-bottom:20px;margin-bottom:28px;display:flex;justify-content:space-between;align-items:flex-end;flex-wrap:wrap;gap:16px}
h1{font-size:30px;font-weight:700;letter-spacing:-.02em}
.tag{font-family:var(--mono);font-size:11px;text-transform:uppercase;letter-spacing:.14em;color:var(--muted);margin-top:6px}
.rundate{font-family:var(--mono);font-size:14px;text-align:right}
.rundate .lbl{font-size:10px;text-transform:uppercase;letter-spacing:.12em;color:var(--muted);display:block}
.sumgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));border:2px solid var(--line);margin-bottom:28px}
.card{padding:18px 20px;border-right:1.5px solid var(--line)}
.card:last-child{border-right:none}
.card .lbl{font-family:var(--mono);font-size:10px;text-transform:uppercase;letter-spacing:.12em;color:var(--muted)}
.card .val{font-size:34px;font-weight:700;letter-spacing:-.02em;margin-top:4px;font-variant-numeric:tabular-nums}
.card.matched .val{color:var(--ok)}.card.cleared .val{color:var(--info)}.card.exceptions .val{color:var(--warn)}
h3{font-family:var(--mono);font-size:11px;text-transform:uppercase;letter-spacing:.14em;color:var(--muted);margin:26px 0 10px;padding-bottom:6px;border-bottom:1.5px solid var(--line)}
table{width:100%;border-collapse:collapse;font-size:13px}
th{font-family:var(--mono);font-size:10px;text-transform:uppercase;letter-spacing:.08em;color:var(--muted);text-align:left;padding:8px 10px;border-bottom:1.5px solid var(--line)}
td{padding:8px 10px;border-bottom:1px solid var(--grid);font-variant-numeric:tabular-nums}
tr:hover td{background:rgba(90,122,82,.05)}
.num{text-align:right}
.pill{font-family:var(--mono);font-size:11px;padding:2px 8px;display:inline-block}
.pill.clean{color:var(--ok);border:1px solid var(--ok)}
.pill.review{color:var(--warn);border:1px solid var(--warn)}
.pill.aged{color:var(--info);border:1px solid var(--info)}
.method{font-family:var(--mono);font-size:11px;color:var(--slate)}
.desc{color:var(--muted);font-size:12px}
.tabs{display:flex;margin:30px 0 0;border-bottom:2px solid var(--line)}
.tab{font-family:var(--mono);font-size:12px;text-transform:uppercase;letter-spacing:.08em;padding:11px 18px;cursor:pointer;border:1.5px solid transparent;border-bottom:none;color:var(--muted);position:relative;top:2px}
.tab.active{color:var(--ink);border-color:var(--line);background:var(--paper);font-weight:600}
.tab .count{font-size:10px;opacity:.7;margin-left:5px}
.panel{display:none;padding-top:18px}.panel.active{display:block}
.tablewrap{max-height:520px;overflow:auto;border:1.5px solid var(--grid)}
.tablewrap th{position:sticky;top:0;background:var(--paper);z-index:1}
.empty{text-align:center;color:var(--ok);padding:24px}
.foot{margin-top:40px;font-family:var(--mono);font-size:11px;color:var(--muted);border-top:1px solid var(--grid);padding-top:16px;line-height:1.7}
"""

_JS = """
document.querySelectorAll('.tab').forEach(function(t){t.onclick=function(){
document.querySelectorAll('.tab').forEach(function(x){x.classList.remove('active')});
document.querySelectorAll('.panel').forEach(function(x){x.classList.remove('active')});
t.classList.add('active');document.getElementById('panel-'+t.dataset.tab).classList.add('active');};});
"""


def _esc(v):
    return html.escape("" if v is None else str(v))


def _rows(records, cols):
    if not records:
        return f'<tr><td colspan="{len(cols)}" class="empty">Nothing here.</td></tr>'
    out = []
    for r in records:
        d = r.as_dict()
        cells = []
        for c, cls in cols:
            cells.append(f'<td class="{cls}">{_esc(d[c])}</td>')
        out.append("<tr>" + "".join(cells) + "</tr>")
    return "".join(out)


def _dict_rows(rows, cols):
    if not rows:
        return f'<tr><td colspan="{len(cols)}" class="empty">Nothing here.</td></tr>'
    out = []
    for d in rows:
        cells = "".join(f'<td class="{cls}">{_esc(d.get(c, ""))}</td>' for c, cls in cols)
        out.append("<tr>" + cells + "</tr>")
    return "".join(out)


def write_html(result: dict, path: Path, lsum: dict = None):
    exc, cle, mat = result["exceptions"], result["cleared"], result["matched"]
    rate = (round(100 * len(mat) / result["internal_count"])
            if result["internal_count"] else 0)

    prov_rows = "".join(
        f'<tr><td>{p["provider"]}</td><td class="num">{p["matched"]}</td>'
        f'<td class="num">{p["cleared"]}</td><td class="num">{p["exceptions"]}</td>'
        f'<td class="num">{p["matched"]+p["cleared"]+p["exceptions"]}</td>'
        f'<td><span class="pill {"review" if p["exceptions"] else "clean"}">'
        f'{"review" if p["exceptions"] else "clean"}</span></td></tr>'
        for p in result["per_provider"])

    exc_cols = [("Provider", ""), ("Side", ""), ("Internal Ref", ""), ("Provider Ref", ""),
                ("Value Date", ""), ("Amount (Internal)", "num"), ("Amount (Provider)", "num"),
                ("Currency", ""), ("Type", ""), ("Description", "desc")]
    cle_cols = [("Provider", ""), ("Side", ""), ("Internal Ref", ""), ("Provider Ref", ""),
                ("Amount (Internal)", "num"), ("Amount (Provider)", "num"), ("Currency", ""),
                ("Cleared Against", ""), ("Method", "method"), ("Description", "desc")]
    mat_cols = [("Provider", ""), ("Internal Ref", ""), ("Provider Ref", ""), ("Value Date", ""),
                ("Amount (Internal)", "num"), ("Amount (Provider)", "num"), ("Difference", "num"),
                ("Currency", ""), ("Method", "method")]

    open_cols = [("Provider", ""), ("Side", ""), ("Internal Ref", ""), ("Provider Ref", ""),
                 ("First Seen", ""), ("Days Open", "num"), ("Amount", "num"), ("Currency", ""),
                 ("Status", ""), ("Description", "desc")]
    clr_cols = [("Provider", ""), ("Side", ""), ("Provider Ref", ""), ("First Seen", ""),
                ("Cleared Date", ""), ("Days To Clear", "num"), ("Amount", "num"),
                ("Currency", ""), ("Cleared Against", "")]

    def thead(cols):
        return "<tr>" + "".join(
            f'<th class="{cls}">{_esc(c)}</th>' for c, cls in cols) + "</tr>"

    # ledger-aware cards + tabs (when a ledger summary is supplied)
    if lsum is not None:
        cards = (
            f'<div class="card matched"><div class="lbl">Matched today</div><div class="val">{len(mat)}</div></div>'
            f'<div class="card cleared"><div class="lbl">Cleared today</div><div class="val">{lsum["cleared_today"]}</div></div>'
            f'<div class="card"><div class="lbl">New exceptions</div><div class="val">{lsum["new_today"]}</div></div>'
            f'<div class="card exceptions"><div class="lbl">Open (total)</div><div class="val">{lsum["open_total"]}</div></div>'
            f'<div class="card exceptions"><div class="lbl">Manual review</div><div class="val">{lsum["open_manual"]}</div></div>')
        manual_rows = [r for r in lsum["open_rows"] if r.get("Status") == "Manual Review"]
        tabs = (
            f'<div class="tab active" data-tab="open">Open exceptions <span class="count">{lsum["open_total"]}</span></div>'
            f'<div class="tab" data-tab="manual">Manual review <span class="count">{len(manual_rows)}</span></div>'
            f'<div class="tab" data-tab="clr">Cleared today <span class="count">{lsum["cleared_today"]}</span></div>'
            f'<div class="tab" data-tab="mat">Matched today <span class="count">{len(mat)}</span></div>')
        panels = (
            f'<div class="panel active" id="panel-open"><div class="tablewrap"><table>'
            f'<thead>{thead(open_cols)}</thead><tbody>{_dict_rows(lsum["open_rows"], open_cols)}</tbody></table></div></div>'
            f'<div class="panel" id="panel-manual"><div class="tablewrap"><table>'
            f'<thead>{thead(open_cols)}</thead><tbody>{_dict_rows(manual_rows, open_cols)}</tbody></table></div></div>'
            f'<div class="panel" id="panel-clr"><div class="tablewrap"><table>'
            f'<thead>{thead(clr_cols)}</thead><tbody>{_dict_rows(lsum["cleared_today_rows"], clr_cols)}</tbody></table></div></div>'
            f'<div class="panel" id="panel-mat"><div class="tablewrap"><table>'
            f'<thead>{thead(mat_cols)}</thead><tbody>{_rows(mat, mat_cols)}</tbody></table></div></div>')
    else:
        cards = (
            f'<div class="card"><div class="lbl">Internal rows</div><div class="val">{result["internal_count"]}</div></div>'
            f'<div class="card matched"><div class="lbl">Matched</div><div class="val">{len(mat)}</div></div>'
            f'<div class="card cleared"><div class="lbl">Cleared (aged)</div><div class="val">{len(cle)}</div></div>'
            f'<div class="card exceptions"><div class="lbl">Exceptions</div><div class="val">{len(exc)}</div></div>'
            f'<div class="card"><div class="lbl">Match rate</div><div class="val">{rate}%</div></div>')
        tabs = (
            f'<div class="tab active" data-tab="exc">Exceptions <span class="count">{len(exc)}</span></div>'
            f'<div class="tab" data-tab="cle">Cleared <span class="count">{len(cle)}</span></div>'
            f'<div class="tab" data-tab="mat">Matched <span class="count">{len(mat)}</span></div>')
        panels = (
            f'<div class="panel active" id="panel-exc"><div class="tablewrap"><table>'
            f'<thead>{thead(exc_cols)}</thead><tbody>{_rows(exc, exc_cols)}</tbody></table></div></div>'
            f'<div class="panel" id="panel-cle"><div class="tablewrap"><table>'
            f'<thead>{thead(cle_cols)}</thead><tbody>{_rows(cle, cle_cols)}</tbody></table></div></div>'
            f'<div class="panel" id="panel-mat"><div class="tablewrap"><table>'
            f'<thead>{thead(mat_cols)}</thead><tbody>{_rows(mat, mat_cols)}</tbody></table></div></div>')

    prior = ", ".join(result["prior_dates"]) or "none available"
    doc = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Recon Daily — {_esc(result['run_date'])}</title>
<style>{_CSS}</style></head><body>
<div class="wrap">
  <header>
    <div><h1>&#9703; Recon Daily</h1>
      <div class="tag">Internal ledger &times; provider statements &middot; aged reconciliation</div></div>
    <div class="rundate"><span class="lbl">Run date</span>{_esc(result['run_date'])}
      <span class="lbl" style="margin-top:8px">Look-back</span>{result['lookback_days']} days</div>
  </header>

  <div class="sumgrid">{cards}</div>

  <h3>By provider (today)</h3>
  <table><thead><tr><th>Provider</th><th class="num">Matched</th><th class="num">Cleared</th>
    <th class="num">Exceptions</th><th class="num">Total</th><th>Status</th></tr></thead>
    <tbody>{prov_rows}</tbody></table>

  <div class="tabs">{tabs}</div>
  {panels}

  <div class="foot">
    Prior days used for aging: {_esc(prior)}.<br>
    Open exceptions carry forward across days and clear automatically when the counterparty reports them
    within {result['lookback_days']} days; still open beyond that &rarr; <b>Manual review</b>.<br>
    Fee rows (Debit_FEE_VAT, Debit_FEE_REMITTANCE_FEE) excluded.
  </div>
</div>
<script>{_JS}</script></body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")
