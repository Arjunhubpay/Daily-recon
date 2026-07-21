"""Cross-run ledger test: an exception opens one day and is automatically
cleared when the counterparty reports it a later day; an item that never
reconciles ages into Manual Review. Run:  python tests/test_ledger.py
"""

import csv
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402
from recon.__main__ import main  # noqa: E402

INT_H = ["Payment Provider", "Reference", "Ref for Account Owner", "Value Date",
         "Amount", "Currency", "Transaction Type", "Supplementary Details"]
NBF_PRE = [["Bank"], ["Statement"], [""], ["Acct"], ["Period"], [""], [""]]
NBF_H = ["Date", "Value Date", "Description", "Debit", "Credit", "Balance"]


def _w(root, day, name, header, rows, preamble=None):
    d = root / day
    d.mkdir(parents=True, exist_ok=True)
    with open(d / name, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        for p in (preamble or []):
            wr.writerow(p)
        wr.writerow(header)
        wr.writerows(rows)


def _open_keys(master):
    wb = openpyxl.load_workbook(master)
    ws = wb["Open Exceptions"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    out = {}
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        out[d["Provider Ref"]] = d
    return out


def _cleared_keys(master):
    wb = openpyxl.load_workbook(master)
    ws = wb["Cleared Log"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    return {dict(zip(hdr, r))["Provider Ref"]: dict(zip(hdr, r)) for r in rows[1:]}


def run():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        out = root / "Output"
        out.mkdir(parents=True, exist_ok=True)

        # Day 07-10: A and C are NBF breaks (not in provider); B matches.
        _w(root, "2026-07-10", "recon-lines.csv", INT_H, [
            ["NBF", "FT10000AAAAA", "", "2026-07-10", "1000", "AED", "N", ""],
            ["NBF", "FT30000CCCCC", "", "2026-07-10", "3000", "AED", "N", ""],
            ["NBF", "FT20000BBBBB", "", "2026-07-10", "2000", "AED", "N", ""],
        ])
        _w(root, "2026-07-10", "OpTransactionHistory.csv", NBF_H, [
            ["2026-07-10", "2026-07-10", "Outward FT20000BBBBB", "2000", "", "0"],
        ], preamble=NBF_PRE)

        main(["--root", str(root), "--date", "2026-07-10", "--quiet"])
        master = out / "Open_Exceptions.xlsx"
        opens = _open_keys(master)
        assert set(opens) == {"FT10000AAAAA", "FT30000CCCCC"}, set(opens)
        assert opens["FT10000AAAAA"]["Status"] == "Open"

        # Day 07-11: provider now reports A -> A should clear; C still open.
        _w(root, "2026-07-11", "recon-lines.csv", INT_H, [
            ["NBF", "FT10000AAAAA", "", "2026-07-11", "1000", "AED", "N", ""],
        ])
        _w(root, "2026-07-11", "OpTransactionHistory.csv", NBF_H, [
            ["2026-07-11", "2026-07-11", "Outward FT10000AAAAA", "1000", "", "0"],
        ], preamble=NBF_PRE)

        main(["--root", str(root), "--date", "2026-07-11", "--quiet"])
        opens = _open_keys(master)
        cleared = _cleared_keys(master)
        assert "FT10000AAAAA" not in opens, "A should have cleared"
        assert "FT30000CCCCC" in opens, "C should still be open"
        assert "FT10000AAAAA" in cleared
        assert cleared["FT10000AAAAA"]["First Seen"] == "2026-07-10"
        assert cleared["FT10000AAAAA"]["Cleared Date"] == "2026-07-11"
        assert cleared["FT10000AAAAA"]["Days To Clear"] == 1

        # dated snapshot exists in the date folder
        assert (root / "2026-07-11" / "Recon_Summary_2026-07-11.xlsx").exists()

        # Day 07-25: C is 15 days old and still missing -> Manual Review.
        _w(root, "2026-07-25", "recon-lines.csv", INT_H, [
            ["NBF", "FT99999ZZZZZ", "", "2026-07-25", "9", "AED", "N", ""],
        ])
        _w(root, "2026-07-25", "OpTransactionHistory.csv", NBF_H, [
            ["2026-07-25", "2026-07-25", "nothing relevant", "", "", "0"],
        ], preamble=NBF_PRE)

        main(["--root", str(root), "--date", "2026-07-25", "--quiet"])
        opens = _open_keys(master)
        assert "FT30000CCCCC" in opens
        assert opens["FT30000CCCCC"]["Status"] == "Manual Review", opens["FT30000CCCCC"]["Status"]
        assert opens["FT30000CCCCC"]["Days Open"] == 15, opens["FT30000CCCCC"]["Days Open"]

    print("test_ledger: OK")


def test_ledger():
    run()


if __name__ == "__main__":
    run()
